import os
import csv
import time
import logging
import urllib.request
import xml.etree.ElementTree as ET
from abc import ABC, abstractmethod
from typing import List, Dict
from datetime import datetime

from iptv_check.models.check_result import CheckResult
from iptv_check.core.epg_parser import EPGParser
from iptv_check.models.epg_program import EpgChannel

logger = logging.getLogger(__name__)


# fanmingming 全量节目单（XMLTV）。与 build_playlist.py 保持一致。
_FANMINGMING_EPG_URL = "https://cdn.jsdelivr.net/gh/fanmingming/live@main/e.xml"

# 全量 EPG 抓取结果的内存缓存（单进程内 6 小时刷新一次，避免每次导出都拉几十 MB）。
_EPG_FULL_CACHE: Dict[str, object] = {"data": None, "url": None, "ts": 0.0}
_EPG_CACHE_TTL = 21600


def _esc_attr(value: str) -> str:
    """清洗 M3U 属性值。

    M3U 行内属性（tvg-id/tvg-name/group-title）不是 XML，无需实体转义——
    实体转义反而会让播放器显示字面 `&amp;`。只需：
    1. 移除换行/控制字符，避免破坏 M3U 行结构；
    2. 值内双引号会提前终止 `tvg-name="..."` 属性，统一替换为单引号。
    """
    if not value:
        return ""
    cleaned = value.replace("\r", "").replace("\n", " ").strip()
    return cleaned.replace('"', "'")


def _region_label(ch) -> str:
    """导出表格用的「国家/地区」列值：优先导出前推断好的归属地，其次国家代码转中文名。"""
    v = getattr(ch, "region", "") or ""
    if v:
        return v
    code = (getattr(ch, "country", "") or "").upper().strip()
    if not code or len(code) != 2:
        return ""
    try:
        from iptv_check.core.parser import _COUNTRY_NAME_ZH
        return _COUNTRY_NAME_ZH.get(code, code)
    except Exception:
        return code


class ExportStrategy(ABC):
    @abstractmethod
    def export(self, results: List[CheckResult], path: str, **kwargs) -> str:
        pass


class M3uExporter(ExportStrategy):
    def render(self, results: List[CheckResult], local_isp: str = "未知", epg_url: str = "",
               with_logo: bool = False, name_fn=None) -> str:
        """渲染 M3U 文本。

        `name_fn(ch)` 若提供，则只用于 EXTINF 逗号后的「播放器显示名」（如加国家前缀），
        `tvg-name` 属性始终写原始名，避免破坏 EPG 节目单匹配。
        """
        header = "#EXTM3U"
        if epg_url:
            header += f' x-tvg-url="{_esc_attr(epg_url)}"'
        lines = [
            header,
            f"# 检测时间: {datetime.now():%Y-%m-%d %H:%M:%S}",
            f"# 本地网络: {local_isp}宽带",
            "",
        ]
        for r in results:
            if not r.is_valid:
                continue
            ch = r.channel
            name = ch.tvg_name or ch.name
            display = name_fn(ch) if name_fn else name
            attrs = []
            if ch.tvg_id:
                attrs.append(f'tvg-id="{_esc_attr(ch.tvg_id)}"')
            attrs.append(f'tvg-name="{_esc_attr(name)}"')
            if ch.logo_url and with_logo:
                # 默认不写台标：远程台标会让 Kodi 等播放器开机逐个下载、启动极慢。
                # 需要台标时由调用方显式传 with_logo=True。
                attrs.append(f'tvg-logo="{_esc_attr(ch.logo_url)}"')
            if ch.group:
                attrs.append(f'group-title="{_esc_attr(ch.group)}"')
            lines.append(f"#EXTINF:-1 {' '.join(attrs)},{_esc_attr(display)}")
            lines.append(ch.url)
        return "\n".join(lines) + "\n"

    def export(self, results: List[CheckResult], path: str, **kwargs) -> str:
        content = self.render(
            results,
            kwargs.get("local_isp", "未知"),
            kwargs.get("epg_url", ""),
            kwargs.get("with_logo", False),
            kwargs.get("name_fn"),
        )
        with open(path, "w", encoding="utf-8") as f:
            f.write(content)
        return path


class TxtExporter(ExportStrategy):
    def export(self, results: List[CheckResult], path: str, **kwargs) -> str:
        with open(path, "w", encoding="utf-8") as f:
            for r in results:
                if not r.is_valid:
                    f.write(f"{r.channel.name},{r.channel.url} # 错误: {r.details}\n")
        return path


class CsvExporter(ExportStrategy):
    HEADERS = ["原始序号", "来源文件", "频道名称", "URL", "状态", "延迟(ms)", "速度(KB/s)", "国家/地区", "信息"]

    def export(self, results: List[CheckResult], path: str, **kwargs) -> str:
        with open(path, "w", encoding="utf-8-sig", newline="") as f:
            writer = csv.writer(f)
            writer.writerow(self.HEADERS)
            for r in results:
                row = list(r.to_tree_values())
                # to_tree_values 为 8 列，把「国家/地区」插在「信息」之前
                writer.writerow(row[:7] + [_region_label(r.channel)] + row[7:])
        return path


class ExcelExporter(ExportStrategy):
    HEADERS = ["原始序号", "来源文件", "频道名称", "URL", "状态", "延迟(ms)", "速度(KB/s)", "国家/地区", "信息"]

    def export(self, results: List[CheckResult], path: str, **kwargs) -> str:
        try:
            import openpyxl
            from openpyxl.styles import Font, Alignment, PatternFill
        except ImportError:
            logger.warning("openpyxl 未安装，跳过Excel导出")
            return ""

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "检测结果"

        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        for col_idx, header in enumerate(self.HEADERS, 1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")

        for row_idx, r in enumerate(results, 2):
            values = list(r.to_tree_values())
            # to_tree_values 为 8 列，把「国家/地区」插在「信息」之前
            values = values[:7] + [_region_label(r.channel)] + values[7:]
            for col_idx, val in enumerate(values, 1):
                cell = ws.cell(row_idx, column=col_idx, value=val)
                if val == "有效":
                    cell.font = Font(color="00B050")
                elif val == "无效":
                    cell.font = Font(color="FF0000")

        ws.auto_filter.ref = ws.dimensions
        for col in ws.columns:
            max_length = max(len(str(cell.value or "")) for cell in col)
            ws.column_dimensions[col[0].column_letter].width = min(max_length + 4, 50)

        wb.save(path)
        return path


def _load_full_epg(epg_url: str) -> Dict[str, EpgChannel]:
    """拉取并解析全量 XMLTV，进程内缓存 6 小时。"""
    now = time.time()
    cached = _EPG_FULL_CACHE
    if cached["data"] is not None and cached["url"] == epg_url and (now - cached["ts"]) < _EPG_CACHE_TTL:
        return cached["data"]  # type: ignore[return-value]
    req = urllib.request.Request(epg_url, headers={"User-Agent": "Mozilla/5.0"})
    with urllib.request.urlopen(req, timeout=120) as resp:
        text = resp.read().decode("utf-8", "replace")
    data = EPGParser.parse_xmltv(text)
    cached["data"] = data
    cached["url"] = epg_url
    cached["ts"] = now
    return data


def build_slim_epg_xml(channels, epg_url: str = _FANMINGMING_EPG_URL, epg_data: dict = None) -> str:
    """从全量 XMLTV 中筛出 channels（每组为 (tvg_name, tvg_id, name)）对应的频道与节目，生成瘦身 EPG。

    复用 EPGParser.match_channel_to_epg 做匹配（支持 tvg-id / tvg-name / 频道名大小写不敏感及归一化模糊匹配），
    因此即使导出 m3u 没有 tvg-name，仅靠频道名也能命中。
    `epg_data` 若已提供（如网页已加载的 EPG 源）则直接使用、不再联网抓取；否则按 `epg_url` 抓取全量节目单。
    """
    if epg_data is None:
        epg_data = _load_full_epg(epg_url)
    keep: Dict[str, EpgChannel] = {}
    for tvg_name, tvg_id, name in channels:
        matched = EPGParser.match_channel_to_epg(name, tvg_id or "", tvg_name or "", epg_data)
        if matched and matched.channel_id not in keep:
            keep[matched.channel_id] = matched

    tv = ET.Element("tv")
    tv.set("generator-info-name", "IPTV-Check slim-epg")
    for ch in keep.values():
        ce = ET.SubElement(tv, "channel", {"id": ch.channel_id})
        dn = ET.SubElement(ce, "display-name")
        dn.text = ch.display_name or ch.channel_id
        for p in ch.programs:
            pe = ET.SubElement(
                tv, "programme",
                {"channel": ch.channel_id, "start": p.start, "stop": p.stop},
            )
            t = ET.SubElement(pe, "title")
            t.text = p.title
            if p.desc:
                d = ET.SubElement(pe, "desc")
                d.text = p.desc
            if p.category:
                c = ET.SubElement(pe, "category")
                c.text = p.category
    body = ET.tostring(tv, encoding="unicode")
    return '<?xml version="1.0" encoding="UTF-8"?>\n' + body


class EpgXmlExporter(ExportStrategy):
    def export(self, results: List[CheckResult], path: str, **kwargs) -> str:
        epg_url = kwargs.get("epg_url", _FANMINGMING_EPG_URL)
        channels = [(r.channel.tvg_name, r.channel.tvg_id, r.channel.name) for r in results if r.is_valid]
        xml_text = build_slim_epg_xml(channels, epg_url, kwargs.get("epg_data"))
        with open(path, "w", encoding="utf-8") as f:
            f.write(xml_text)
        return path


class ExportEngine:
    def __init__(self):
        self._strategies = {
            "m3u": M3uExporter(),
            "m3u8": M3uExporter(),
            "txt": TxtExporter(),
            "csv": CsvExporter(),
            "xlsx": ExcelExporter(),
            "epg": EpgXmlExporter(),
        }

    def export(self, fmt: str, results: List[CheckResult], path: str, **kwargs) -> str:
        strategy = self._strategies.get(fmt)
        if not strategy:
            raise ValueError(f"不支持的导出格式: {fmt}")
        return strategy.export(results, path, **kwargs)

    def export_batch(self, formats: List[str], results: List[CheckResult],
                     export_dir: str, base_name: str, local_isp: str = "未知",
                     epg_url: str = "", epg_data=None, with_logo: bool = False,
                     name_fn=None) -> List[str]:
        exported = []
        name_map = {
            "m3u": f"{base_name}_有效源.m3u",
            "m3u8": f"{base_name}_有效源.m3u8",
            "txt": f"{base_name}_无效源.txt",
            "csv": f"{base_name}_检测结果.csv",
            "xlsx": f"{base_name}_检测结果.xlsx",
            "epg": f"{base_name}_节目单.epg.xml",
        }
        for fmt in formats:
            path = os.path.join(export_dir, name_map[fmt])
            kw = {"local_isp": local_isp}
            # 普通 m3u/m3u8 导出也内嵌 x-tvg-url，指向同名的本地瘦身 epg.xml，
            # 这样单独导出 m3u 也能和后续导出的 epg 直接配套（Kodi 按同目录解析）。
            if fmt in ("m3u", "m3u8"):
                kw["epg_url"] = epg_url or name_map["epg"]
                # 台标默认不写（远程台标会让 Kodi 开机逐个下载、启动极慢），
                # 由调用方按需打开 with_logo。
                kw["with_logo"] = with_logo
                # 可选：仅改播放器显示名（tvg-name 保持原名，避免破坏 EPG 匹配）
                kw["name_fn"] = name_fn
            if fmt == "epg":
                kw["epg_data"] = epg_data
            result = self.export(fmt, results, path, **kw)
            if result:
                exported.append(result)
        return exported

    def export_m3u(self, results: List[CheckResult], local_isp: str = "未知") -> str:
        """返回 M3U 文本（供 HTTP 服务直接对外提供），不写文件。"""
        return M3uExporter().render(results, local_isp)
